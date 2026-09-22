# -*- coding: utf-8 -*-
"""
export_proyeccion.py -- Lee la hoja DATA_EXPORT de CADIZ_Gestion_v2.xlsx (el modelo de proyección
financiera de CADIZ, construido y auditado por separado) para el tablero web.

DISEÑO (corregido -- ver Adenda 10 / hotfix README): la web YA NO depende de un CSV intermedio que
haya que regenerar a mano. `read_dataframe()` lee el Excel DIRECTAMENTE cada vez que la app lo
necesita -- el flujo de trabajo del usuario es simplemente reemplazar `CADIZ_Gestion_v2.xlsx` en la
raíz del repo (siempre el mismo nombre de archivo) cada vez que haya una versión nueva; la próxima
carga de la app toma automáticamente esa versión como "la más actual". No hace falta correr ningún
script a mano ni subir ningún CSV.

DATA_EXPORT ya trae TODAS las rondas (históricas REAL y proyectadas PLAN/SIM de CADIZ) en una sola
tabla larga (round, scenario, status, team, region, technology, metric, value, unit, source_value,
source_unit, is_canonical), así que una sola lectura del Excel alcanza -- no hace falta versionar
nada por ronda.

Filtra SOLO team == 'CADIZ' -- el resto de los equipos (dato de competidores) ya le llega a la web
por su propio camino (los .xls de RDOS parseados por cesim_parser.py); DATA_EXPORT es la fuente de
la posición y la proyección PROPIA de CADIZ, no de la competencia.

Este módulo también se puede seguir usando desde la línea de comandos para volcar un CSV de
diagnóstico/respaldo puntual (`export()`), pero eso ya NO es parte del flujo normal de actualización
de la web -- ver `gap_analysis.load_proyeccion()`, que llama a `read_dataframe()` directamente.

Uso (opcional, solo para un CSV de respaldo manual):
    python3 export_proyeccion.py [ruta_al_excel] [ruta_salida_csv]
"""
import csv
import sys
import os

import openpyxl
import pandas as pd

DATA_EXPORT_SHEET = "DATA_EXPORT"
COLUMNS = ["round", "scenario", "status", "team", "region", "technology", "metric", "value", "unit",
           "source_value", "source_unit", "is_canonical"]

# CADIZ_Gestion_v2.xlsx vive siempre en la raíz del repo, al lado de este script y de app.py --
# el usuario reemplaza este archivo (mismo nombre) cada vez que hay una versión nueva del modelo.
DEFAULT_EXCEL = os.path.join(os.path.dirname(os.path.abspath(__file__)), "CADIZ_Gestion_v2.xlsx")
DEFAULT_OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "data", "proyeccion", "cadiz_proyeccion.csv")


def read_dataframe(excel_path=DEFAULT_EXCEL, team="CADIZ") -> pd.DataFrame:
    """Lee DATA_EXPORT de excel_path y devuelve un DataFrame filtrado a `team`, con las columnas del
    contrato (COLUMNS). Lanza ValueError si falta la hoja o si el encabezado no coincide con lo
    esperado (para no desalinear el crosswalk en silencio) -- el llamador (gap_analysis.load_proyeccion)
    es responsable de capturar esto y degradar con gracia (la web debe poder arrancar igual)."""
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    if DATA_EXPORT_SHEET not in wb.sheetnames:
        raise ValueError(f"No se encontró la hoja '{DATA_EXPORT_SHEET}' en {excel_path}")
    ws = wb[DATA_EXPORT_SHEET]

    header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    if header != COLUMNS:
        raise ValueError(
            f"El encabezado de {DATA_EXPORT_SHEET} cambió respecto al contrato esperado.\n"
            f"Esperado: {COLUMNS}\nEncontrado: {header}\n"
            "Revisar el crosswalk (metric_crosswalk.py) antes de usar este Excel -- puede haber "
            "quedado desalineado con nombres de columna nuevos."
        )

    rows = [row for row in ws.iter_rows(min_row=2, values_only=True) if row[3] == team]
    return pd.DataFrame(rows, columns=COLUMNS)


def export(excel_path=DEFAULT_EXCEL, out_path=DEFAULT_OUT, team="CADIZ"):
    """Vuelca a CSV -- utilidad de respaldo/diagnóstico manual, YA NO es parte del flujo normal de
    actualización de la web (ver docstring del módulo)."""
    df = read_dataframe(excel_path, team=team)
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    df.to_csv(out_path, index=False, quoting=csv.QUOTE_MINIMAL)

    print(f"[export_proyeccion] {excel_path} -> {out_path}")
    print(f"  Filas exportadas (team={team}): {len(df)}")
    print(f"  Rondas presentes: {sorted(r for r in df['round'].unique() if r is not None)}")
    print(f"  Status presentes: {sorted(df['status'].unique())}")
    return len(df)


if __name__ == "__main__":
    excel_arg = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_EXCEL
    out_arg = sys.argv[2] if len(sys.argv) > 2 else DEFAULT_OUT
    export(excel_arg, out_arg)
