# -*- coding: utf-8 -*-
"""Estilos minimos y consistentes para CADIZ_Gestion_v1.xlsx (autocontenido, sin dependencia del
proyecto anterior). Paleta simple: azul=input editable, negro=formula/calculo, verde=link a otra
hoja, gris=historico/constante, rojo=alerta/control, texto en Arial en todo el workbook."""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FONT_NAME = "Arial"

COLOR_INPUT = "0000FF"       # azul -- decision/condicion editable
COLOR_CALC = "000000"        # negro -- formula interna
COLOR_LINK = "006100"        # verde -- link a otra hoja
COLOR_HIST = "595959"        # gris -- historico/constante estructural
COLOR_ALERT = "C00000"       # rojo -- alerta/control
COLOR_OUTPUT = "000000"

FILL_HEADER = PatternFill("solid", fgColor="1F4E78")
FILL_SECTION = PatternFill("solid", fgColor="D9E1F2")
FILL_TITLE = PatternFill("solid", fgColor="1F4E78")
FILL_INPUT_BG = PatternFill("solid", fgColor="FFF2CC")
FILL_OK = PatternFill("solid", fgColor="C6EFCE")
FILL_WARN = PatternFill("solid", fgColor="FFEB9C")
FILL_ERROR = PatternFill("solid", fgColor="FFC7CE")

THIN = Side(style="thin", color="BFBFBF")
BORDER_THIN = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_LEFT_WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True)
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")

NUM_MONEY = '#,##0;(#,##0);"-"'
NUM_UNITS = '#,##0;(#,##0);"-"'
NUM_PRICE = '#,##0.00;(#,##0.00);"-"'
NUM_PCT1 = '0.0%;(0.0%);"-"'
NUM_PCT2 = '0.00%;(0.00%);"-"'
NUM_X = '0.00"x"'

KIND_COLOR = {
    "input": COLOR_INPUT, "calculo": COLOR_CALC, "link": COLOR_LINK,
    "historico": COLOR_HIST, "alerta": COLOR_ALERT, "output": COLOR_OUTPUT,
    "plain": COLOR_CALC,
}
KIND_FILL = {
    "input": FILL_INPUT_BG,
}


def f_normal(size=10, bold=False, italic=False, color="000000"):
    return Font(name=FONT_NAME, size=size, bold=bold, italic=italic, color=color)


def set_col_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def style_title_row(ws, row, text, ncols):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    c = ws.cell(row=row, column=1, value=text)
    c.font = f_normal(13, bold=True, color="FFFFFF")
    c.fill = FILL_TITLE
    c.alignment = ALIGN_LEFT
    for col in range(1, ncols + 1):
        ws.cell(row=row, column=col).fill = FILL_TITLE


def style_header_row(ws, row, headers):
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = f_normal(9, bold=True, color="FFFFFF")
        c.fill = FILL_HEADER
        c.alignment = ALIGN_CENTER
        c.border = BORDER_THIN


def style_section_row(ws, row, text, ncols):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
    c = ws.cell(row=row, column=1, value=text)
    c.font = f_normal(10, bold=True, color="1F4E78")
    c.fill = FILL_SECTION
    c.alignment = ALIGN_LEFT
    for col in range(1, ncols + 1):
        ws.cell(row=row, column=col).fill = FILL_SECTION


def apply_cell(ws, row, col, value=None, kind="calculo", numfmt=None, align=None,
               bold=False, italic=False, size=10):
    c = ws.cell(row=row, column=col, value=value)
    color = KIND_COLOR.get(kind, COLOR_CALC)
    c.font = f_normal(size, bold=bold, italic=italic, color=color)
    if numfmt:
        c.number_format = numfmt
    if align:
        c.alignment = align
    fill = KIND_FILL.get(kind)
    if fill:
        c.fill = fill
    return c
